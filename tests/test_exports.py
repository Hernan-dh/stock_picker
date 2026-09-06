import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
from zipfile import ZipFile

from docx import Document
import gradio as gr
from report_export import write_report, report_blocks, prepare_with_downloads, finish_with_downloads


REPORT = """# Informe financiero

Investigación: España, año 2026. Ingresos: €150 y variación −2%.

## Fuentes

- [Fuente](https://example.com/report?a=1&b=2)

| Empresa | Resultado |
| --- | --- |
| Ejemplo | €150 |
"""


class ExportTests(unittest.TestCase):
    def test_markdown_preserves_original_exactly(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_report(REPORT, "md", Path(directory))
            self.assertEqual(path.read_text(encoding="utf-8"), REPORT)

    def test_docx_contains_headings_sources_and_table(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_report(REPORT, "docx", Path(directory))
            with ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
            doc = Document(path)
            self.assertEqual(doc.paragraphs[0].text, "Informe financiero")
            self.assertEqual(doc.paragraphs[0].style.name, "Heading 1")
            self.assertIn("https://example.com/report?a=1&b=2", "\n".join(p.text for p in doc.paragraphs))
            self.assertEqual(doc.tables[0].cell(1, 1).text, "€150")

    def test_pdf_handles_unicode_long_reports_and_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_report(REPORT * 50, "pdf", Path(directory))
            content = path.read_bytes()
            self.assertTrue(content.startswith(b"%PDF-"))
            self.assertTrue(content.rstrip().endswith(b"%%EOF"))
            self.assertGreater(len(content), 1000)

    def test_no_empty_reports_or_arbitrary_formats(self):
        with tempfile.TemporaryDirectory() as directory:
            for report, extension in [("", "md"), (REPORT, "../other"), (None, "pdf")]:
                with self.subTest(extension=extension), self.assertRaises(ValueError):
                    write_report(report, extension, Path(directory))
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_literal_html_is_not_executed(self):
        blocks = report_blocks('<script>alert("test")</script>')
        self.assertIn("<script>", str(blocks))

    def test_preparation_clears_previous_download_state(self):
        submit = Mock(return_value=("", [{"role": "user", "content": "Example"}], gr.Button()))
        result = prepare_with_downloads(submit)("Example", [])
        self.assertIsNone(result[3])
        self.assertFalse(result[4].visible)

    def test_only_final_result_is_exportable(self):
        for text, expected in [("Final report", "Final report"), ("Failed", None)]:
            finish = Mock(return_value=(gr.Textbox(), [
                {"role": "user", "content": "Example"}, {"role": "assistant", "content": text},
            ], gr.Button()))
            result = finish_with_downloads(finish, "Failed")([])
            self.assertEqual(result[3], expected)

    def test_files_from_different_requests_are_isolated(self):
        with tempfile.TemporaryDirectory() as first, tempfile.TemporaryDirectory() as second:
            a = write_report("First", "md", Path(first))
            b = write_report("Second", "md", Path(second))
            self.assertNotEqual(a, b)
            self.assertEqual(a.read_text(), "First")
            self.assertEqual(b.read_text(), "Second")
