"""Report export and session-local Gradio download controls."""
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.sax.saxutils import escape

import gradio as gr
from docx import Document
from markdown_it import MarkdownIt
import reportlab
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer

FONT = "ReportVera"
pdfmetrics.registerFont(TTFont(FONT, str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")))


def inline_text(token):
    """Keep readable text and link destinations without loading external assets."""
    parts, links = [], []
    for child in token.children or []:
        if child.type == "link_open":
            links.append(child.attrGet("href") or "")
        elif child.type == "link_close":
            url = links.pop() if links else ""
            if url:
                parts.append(f" ({url})")
        elif child.type in {"softbreak", "hardbreak"}:
            parts.append("\n")
        elif child.type in {"text", "code_inline", "image"}:
            parts.append(child.content)
    return "".join(parts)


def report_blocks(markdown):
    """Reduce Markdown to portable headings, paragraphs, lists, code and tables."""
    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(markdown)
    blocks, lists = [], []
    kind, heading, rows, row = "paragraph", 0, None, None
    for token in tokens:
        if token.type == "table_open":
            rows = []
        elif token.type == "tr_open":
            row = []
        elif token.type == "tr_close":
            rows.append(row)
        elif token.type == "table_close":
            blocks.append(("table", rows))
            rows = None
        elif token.type == "heading_open":
            kind, heading = "heading", int(token.tag[1])
        elif token.type == "paragraph_open":
            kind = "paragraph"
        elif token.type in {"bullet_list_open", "ordered_list_open"}:
            lists.append(None if token.type == "bullet_list_open" else int(token.attrGet("start") or 1))
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "inline":
            text = inline_text(token)
            if rows is not None:
                row.append(text)
            elif kind == "heading":
                blocks.append(("heading", (heading, text)))
                kind = "paragraph"
            else:
                if lists:
                    marker = lists[-1]
                    text = ("• " if marker is None else f"{marker}. ") + text
                    if marker is not None:
                        lists[-1] += 1
                blocks.append(("paragraph", text))
        elif token.type in {"fence", "code_block"}:
            blocks.append(("code", token.content))
    return blocks


def write_report(markdown: str, file_format: str, directory: Path) -> Path:
    """Export one completed report; filenames never depend on user input."""
    if not isinstance(markdown, str) or not markdown.strip():
        raise ValueError("A completed report is required")
    if file_format not in {"md", "docx", "pdf"}:
        raise ValueError("Unsupported export format")
    path = directory / f"report.{file_format}"
    if file_format == "md":
        path.write_text(markdown, encoding="utf-8")
        return path
    blocks = report_blocks(markdown)
    if file_format == "docx":
        document = Document()
        for kind, value in blocks:
            if kind == "heading":
                document.add_heading(value[1], level=min(value[0], 9))
            elif kind == "table":
                table = document.add_table(rows=0, cols=max(map(len, value)))
                table.style = "Table Grid"
                for cells in value:
                    for cell, text in zip(table.add_row().cells, cells):
                        cell.text = text
            else:
                document.add_paragraph(value)
        document.save(path)
    else:
        style = ParagraphStyle("Report", fontName=FONT, fontSize=10, leading=15, spaceAfter=8)
        document = SimpleDocTemplate(str(path), rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
        def paragraph(text, selected_style=style):
            return Paragraph(escape(text).replace("\n", "<br/>"), selected_style)
        story = []
        for kind, value in blocks:
            if kind == "heading":
                level, text = value
                heading_style = ParagraphStyle(
                    f"Heading{level}", parent=style, fontSize=max(11, 20 - level * 2),
                    leading=23, spaceBefore=10, keepWithNext=True,
                )
                story.append(paragraph(text, heading_style))
            elif kind == "table":
                columns = max(map(len, value))
                cells = [[paragraph(text) for text in row] + [""] * (columns - len(row)) for row in value]
                story.append(LongTable(
                    cells, colWidths=[document.width / columns] * columns,
                    repeatRows=1, splitInRow=1,
                    style=[("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")],
                ))
                story.append(Spacer(1, 8))
            else:
                story.append(paragraph(value))
        document.build(story or [paragraph(markdown)])
    return path


def empty_download():
    return gr.File(value=None, visible=False)


def download_controls(language):
    """Create controls and a private report state for one localized chat."""
    spanish = language == "Español"
    report = gr.State(None)
    with gr.Row(elem_id="download-controls-es" if spanish else "download-controls-en"):
        gr.Markdown("Formato" if spanish else "Format", elem_classes="format-label")
        file_format = gr.Dropdown(
            choices=[("Markdown (.md)", "md"), ("Word (.docx)", "docx"), ("PDF (.pdf)", "pdf")],
            value="md", show_label=False, interactive=True, scale=1,
        )
        button = gr.Button("Preparar descarga" if spanish else "Prepare download", scale=2)
    output = gr.File(label="Descargar informe" if spanish else "Download report", visible=False, interactive=False)

    def download(markdown, selected_format):
        if not markdown:
            raise gr.Error("Primero completá una consulta." if spanish else "Complete a request first.")
        try:
            with TemporaryDirectory(prefix="report-export-") as directory:
                path = write_report(markdown, selected_format, Path(directory))
                cached = output.move_resource_to_block_cache(path)
            return gr.File(value=cached, visible=True)
        except Exception as error:
            print(f"[export] failed ({type(error).__name__})", flush=True)
            raise gr.Error(
                "No pude generar el archivo. Intentá con Markdown." if spanish
                else "Could not generate the file. Please try Markdown."
            ) from None

    button.click(download, [report, file_format], output, show_progress="hidden")
    file_format.change(empty_download, outputs=output, queue=False)
    return report, output


def prepare_with_downloads(submit):
    def prepare(message, history):
        return (*submit(message, history), None, empty_download())
    return prepare


def finish_with_downloads(finish, error_text):
    def complete(history):
        textbox, updated, button = finish(history)
        report = updated[-1]["content"] if len(updated) >= 2 else None
        if report == error_text:
            report = None
        return textbox, updated, button, report, empty_download()
    return complete


def finish_with_progress_downloads(finish, error_text):
    """Preserve report downloads while a Gradio generator publishes task progress."""
    def complete(history):
        for textbox, updated, button, completed in finish(history):
            report = updated[-1]["content"] if len(updated) >= 2 else None
            if not completed or report == error_text:
                report = None
            yield textbox, updated, button, report, empty_download()
    return complete
